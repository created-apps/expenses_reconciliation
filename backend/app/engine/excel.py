"""The client's Excel workbook (project_context.md section 7), built from plain values.

Pure: no database, network or clock. The service passes in the lines, their invoice file names, the unused
invoices and the statement's printed figures; the workbook comes back as bytes.

Sheets, in order: `Summary`, `<Month YYYY>`, `Invoices not used`. Arial 10 throughout. The main sheet's
columns A to E are the client's format: Date, Transaction ID, Description, Amount, File Name. Money is kept as
integer paise until the moment it is written as a number (a Decimal, never a float).
"""

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.engine.money import short_decimal

STATUSES = (
    "Matched", "Review", "Missing invoice", "No invoice", "Reversed", "Credit – manual", "Payment", "Bank charge",
)
HEADERS = ("Date", "Transaction ID", "Description", "Amount", "File Name", "Status", "Notes")
UNUSED_HEADERS = ("File", "Vendor", "Date", "Amount", "Reason")

DATE_FORMAT = "d/m/yyyy"
INR_FORMAT = '"₹"#,##0.00;-"₹"#,##0.00'  # the client's format; the symbol is quoted so every Excel reads it
USD_FORMAT = '"$"#,##0.00;-"$"#,##0.00'
PLAIN_FORMAT = "#,##0.00;-#,##0.00"

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
RULE = Border(bottom=Side(style="thin", color="999999"))

MONTHS = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)


# -- inputs -----------------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class SheetLine:
    txn_date: date
    ser_no: str
    description: str
    amount_paise: int  # always positive; is_credit marks credits
    is_credit: bool
    foreign_amount_minor: int | None
    foreign_currency: str | None
    files: tuple[str, ...]  # invoice file names (attached and suggested), in the order to print
    status: str | None
    note: str


@dataclass(frozen=True)
class UnusedRow:
    file_name: str
    vendor: str  # blank when no vendor is set
    when: date | None
    currency: str
    total_minor: int | None
    reason: str


@dataclass(frozen=True)
class Statement:
    """The figures the statement prints on page 1 (paise); None when they were not read."""

    debits_paise: int | None  # Purchases/Charges plus Cash Advances
    credits_paise: int | None  # Payments/Credits
    period_start: date | None = None
    period_end: date | None = None


@dataclass(frozen=True)
class TotalsCheck:
    debits_paise: int
    credits_paise: int
    printed_debits_paise: int | None
    printed_credits_paise: int | None
    debits_difference_paise: int | None  # lines minus printed
    credits_difference_paise: int | None
    ok: bool


# -- small helpers -------------------------------------------------------------------------------------------------
def money(paise: int) -> Decimal:
    """Integer paise (or cents) as the number to write in a cell: exact, never a float."""
    return Decimal(paise) / 100


def check_totals(lines: list[SheetLine], statement: Statement) -> TotalsCheck:
    """Debits and credits as the sheet adds them up, against what the statement prints."""
    debits = sum(x.amount_paise for x in lines if not x.is_credit)
    credits_ = sum(x.amount_paise for x in lines if x.is_credit)
    d = None if statement.debits_paise is None else debits - statement.debits_paise
    c = None if statement.credits_paise is None else credits_ - statement.credits_paise
    return TotalsCheck(debits, credits_, statement.debits_paise, statement.credits_paise, d, c, d == 0 and c == 0)


def sheet_name_for(statement_date: date | None, label: str | None) -> str:
    """`June 2026` from the statement date; the cycle label if there is no date; never an illegal sheet name."""
    if statement_date:
        return f"{MONTHS[statement_date.month - 1]} {statement_date.year}"
    cleaned = "".join(ch for ch in (label or "") if ch not in "[]:*?/\\'").strip()[:31]
    return cleaned or "Statement"


def export_file_name(sheet_name: str) -> str:
    """`CC Expense June 2026.xlsx`, after the client's own file name."""
    cleaned = "".join(ch for ch in sheet_name if ch not in '<>:"/\\|?*').strip()
    return f"CC Expense {cleaned}.xlsx" if cleaned else "CC Expense.xlsx"


def line_description(description: str, foreign_amount_minor: int | None, foreign_currency: str | None) -> str:
    """Foreign lines append `  [86.12 USD]`; whole or round amounts are written short, as in the client's
    earlier workbooks (`[36 USD]`, `[133.5 USD]`)."""
    if foreign_amount_minor is None or not foreign_currency:
        return description
    return f"{description}  [{short_decimal(foreign_amount_minor)} {foreign_currency}]"


def _put(ws, row: int, col: int, value, *, bold: bool = False, fmt: str | None = None, align: str | None = None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BOLD if bold else FONT
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = Alignment(horizontal=align)
    return cell


def _header(ws, headers: tuple[str, ...]) -> None:
    for i, text in enumerate(headers, start=1):
        cell = _put(ws, 1, i, text, bold=True)
        cell.border = RULE
    ws.freeze_panes = "A2"


def _widths(ws, widths: tuple[int, ...]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _unused_format(currency: str) -> str:
    return INR_FORMAT if currency == "INR" else USD_FORMAT if currency == "USD" else PLAIN_FORMAT


# -- the workbook ----------------------------------------------------------------------------------------------------
def build_workbook(sheet_name: str, lines: list[SheetLine], unused: list[UnusedRow], statement: Statement) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    main = wb.create_sheet(sheet_name)
    extra = wb.create_sheet("Invoices not used")

    last = len(lines) + 1  # last row of the main sheet (the header is row 1)
    _fill_main(main, lines)
    _fill_unused(extra, unused)
    _fill_summary(summary, sheet_name, last, statement)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _fill_main(ws, lines: list[SheetLine]) -> None:
    _header(ws, HEADERS)
    for r, x in enumerate(lines, start=2):
        _put(ws, r, 1, x.txn_date, fmt=DATE_FORMAT, align="left")
        # The serial is a number written as `0` so Excel shows all its digits (13598184078), not 1.36E+10.
        _put(ws, r, 2, int(x.ser_no) if x.ser_no.isdigit() else x.ser_no, fmt="0", align="left")
        _put(ws, r, 3, line_description(x.description, x.foreign_amount_minor, x.foreign_currency))
        _put(ws, r, 4, money(-x.amount_paise if x.is_credit else x.amount_paise), fmt=INR_FORMAT)
        _put(ws, r, 5, ", ".join(x.files))
        _put(ws, r, 6, x.status or "")
        _put(ws, r, 7, x.note)
    _widths(ws, (12, 16, 46, 16, 60, 16, 70))
    if lines:
        ws.auto_filter.ref = f"A1:G{len(lines) + 1}"


def _fill_unused(ws, rows: list[UnusedRow]) -> None:
    _header(ws, UNUSED_HEADERS)
    for r, x in enumerate(rows, start=2):
        _put(ws, r, 1, x.file_name)
        _put(ws, r, 2, x.vendor)
        _put(ws, r, 3, x.when, fmt=DATE_FORMAT, align="left")
        _put(ws, r, 4, None if x.total_minor is None else money(x.total_minor), fmt=_unused_format(x.currency))
        _put(ws, r, 5, x.reason)
    _widths(ws, (44, 22, 12, 16, 70))


def _fill_summary(ws, sheet_name: str, last: int, statement: Statement) -> None:
    """Counts and amounts per status are formulas over the main sheet, so they follow any edit made in Excel."""
    quoted = f"'{sheet_name}'"
    status_range = f"{quoted}!$F$2:$F${max(last, 2)}"
    amount_range = f"{quoted}!$D$2:$D${max(last, 2)}"

    _put(ws, 1, 1, "Reconciliation summary", bold=True)
    if statement.period_start and statement.period_end:
        _put(ws, 2, 1, f"Statement period {statement.period_start:%d/%m/%Y} to {statement.period_end:%d/%m/%Y}")

    for col, text in enumerate(("Status", "Lines", "Amount"), start=1):
        _put(ws, 4, col, text, bold=True).border = RULE
    first = 5
    for i, name in enumerate(STATUSES):
        r = first + i
        _put(ws, r, 1, name)
        _put(ws, r, 2, f"=COUNTIF({status_range},A{r})", fmt="0")
        _put(ws, r, 3, f"=SUMIF({status_range},A{r},{amount_range})", fmt=INR_FORMAT)
    end = first + len(STATUSES) - 1
    total = end + 1
    _put(ws, total, 1, "Total", bold=True)
    _put(ws, total, 2, f"=SUM(B{first}:B{end})", bold=True, fmt="0")
    _put(ws, total, 3, f"=SUM(C{first}:C{end})", bold=True, fmt=INR_FORMAT)

    block = total + 2
    _put(ws, block, 1, "Check against the statement", bold=True)
    for col, text in enumerate(("", "Sheet", "Statement", "Difference"), start=1):
        _put(ws, block + 1, col, text, bold=True).border = RULE
    rows = (
        ("Debits (Purchases/Charges)", f'=SUMIF({amount_range},">0")', statement.debits_paise),
        ("Credits (Payments/Credits)", f'=-SUMIF({amount_range},"<0")', statement.credits_paise),
    )
    for k, (label, formula, printed) in enumerate(rows):
        r = block + 2 + k
        _put(ws, r, 1, label)
        _put(ws, r, 2, formula, fmt=INR_FORMAT)
        if printed is None:
            _put(ws, r, 3, "not read")
            _put(ws, r, 4, "n/a")
        else:
            _put(ws, r, 3, money(printed), fmt=INR_FORMAT)
            _put(ws, r, 4, f"=ROUND(B{r}-C{r},2)", fmt=INR_FORMAT)  # rounded: sums of decimals can differ at the 10th place
    d, c = block + 2, block + 3
    _put(
        ws, block + 4, 1,
        f'=IF(AND(D{d}=0,D{c}=0),"OK: the sheet adds up to the statement","Does not add up to the statement")',
        bold=True,
    )
    _widths(ws, (34, 22, 22, 16))
