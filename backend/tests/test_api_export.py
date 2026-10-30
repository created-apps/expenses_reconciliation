"""GET /export-summary and /export.xlsx, and what an edit after an export does (fakes for the database).

Uses the same six-line reconciliation as the Slice 10 tests (a Review with a suggestion, Missing invoices,
a USD charge, a payment, unused invoices).
"""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import CurrentUser, get_current_user
from app.config import Settings, get_settings
from app.db import get_repo
from app.main import app
from fakes import FakeRepo
from test_api_results import World

USER = CurrentUser(email="harsh@create-ed.in", user_id="u1")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DEBITS = 42200 + 5253900 + 71501 + 354437 + 90000  # the five debit lines of the world
CREDITS = 6047200


@pytest.fixture
def api():
    repo = FakeRepo()
    app.dependency_overrides[get_current_user] = lambda: USER
    app.dependency_overrides[get_repo] = lambda: repo
    app.dependency_overrides[get_settings] = lambda: Settings(_env_file=None)
    client = TestClient(app)
    world = World(repo)
    repo.update_cycle(world.cid, {"summary_purchases_paise": DEBITS, "summary_credits_paise": CREDITS,
                                 "period_start": "2026-05-21", "period_end": "2026-06-28"})
    client.world = world
    yield client
    app.dependency_overrides.clear()


def summary(api, expected=200):
    r = api.get(f"/api/cycles/{api.world.cid}/export-summary")
    assert r.status_code == expected, r.text
    return r.json()


def download(api, expected=200):
    r = api.get(f"/api/cycles/{api.world.cid}/export.xlsx")
    assert r.status_code == expected, r.text
    return r


def cycle(api):
    return api.get(f"/api/cycles/{api.world.cid}").json()


def sheet_rows(response, name):
    wb = load_workbook(BytesIO(response.content))
    return wb, [[c.value for c in row] for row in wb[name].iter_rows()]


# ---- gates ----------------------------------------------------------------------------------------------------------------------------
def test_the_endpoints_need_a_token():
    client = TestClient(app)
    nil = "6f9619ff-8b86-d011-b42d-00c04fc964ff"
    assert client.get(f"/api/cycles/{nil}/export-summary").status_code == 401
    assert client.get(f"/api/cycles/{nil}/export.xlsx").status_code == 401


def test_unknown_cycle_is_404(api):
    nil = "6f9619ff-8b86-d011-b42d-00c04fc964ff"
    assert api.get(f"/api/cycles/{nil}/export-summary").status_code == 404
    assert api.get(f"/api/cycles/{nil}/export.xlsx").status_code == 404


def test_nothing_can_be_exported_before_matching_has_run(api):
    api.world.row("L3")["status"] = None
    assert "Run matching" in summary(api, 409)["detail"]
    assert "Run matching" in download(api, 409).json()["detail"]
    assert cycle(api)["status"] == "matched" and not [e for e in api.world.repo.events if e["action"] == "exported"]


def test_a_cycle_whose_statement_is_not_confirmed_cannot_be_exported(api):
    api.world.repo.update_cycle(api.world.cid, {"status": "statement_parsed"})
    assert download(api, 409).json()["detail"] == "Confirm the statement first"


def test_nothing_can_be_exported_while_matching_runs(api):
    w = api.world
    w.repo.update_cycle(w.cid, {"job_kind": "matching", "job_status": "running", "job_progress": 30, "job_message": "Matching 6 lines…"})
    assert "Matching is running" in download(api, 409).json()["detail"]


# ---- the summary the Export step shows ------------------------------------------------------------------------------------------------
def test_the_summary_has_the_counts_the_totals_check_and_the_review_warning(api):
    s = summary(api)
    assert (s["sheet_name"], s["file_name"]) == ("June 2026", "CC Expense June 2026.xlsx")
    assert s["sheets"] == ["Summary", s["sheet_name"], "Invoices not used"]
    assert s["lines"] == 6 and s["needs_review"] == 3  # one Review and two Missing invoice
    counts = {x["status"]: x["count"] for x in s["statuses"]}
    assert counts["Review"] == 1 and counts["Missing invoice"] == 2 and counts["Matched"] == 2 and counts["Payment"] == 1
    assert (s["debits_paise"], s["printed_debits_paise"], s["debits_difference_paise"]) == (DEBITS, DEBITS, 0)
    assert (s["credits_paise"], s["printed_credits_paise"], s["credits_difference_paise"]) == (CREDITS, CREDITS, 0)
    assert s["totals_ok"] is True and s["unused_invoices"] == 5
    assert s["last_exported_at"] is None and s["edited_since_export"] is False


def test_the_summary_shows_a_difference_when_the_lines_do_not_add_up_to_the_statement(api):
    w = api.world
    w.repo.update_cycle(w.cid, {"summary_purchases_paise": DEBITS + 1234})
    s = summary(api)
    assert s["debits_difference_paise"] == -1234 and s["totals_ok"] is False


def test_the_month_sheet_is_named_after_the_statement_date(api):
    w = api.world
    assert summary(api)["sheet_name"] == "June 2026"  # statement_date is 28 June 2026 in the world
    w.repo.update_cycle(w.cid, {"statement_date": "2026-09-30"})
    assert summary(api)["file_name"] == "CC Expense September 2026.xlsx"


# ---- the download -----------------------------------------------------------------------------------------------------------------------
def test_the_download_is_the_workbook_of_the_current_state(api):
    w = api.world
    r = download(api)
    assert r.headers["content-type"] == XLSX
    assert "CC%20Expense%20June%202026.xlsx" in r.headers["content-disposition"]
    wb, rows = sheet_rows(r, "June 2026")
    assert wb.sheetnames == ["Summary", "June 2026", "Invoices not used"]
    assert rows[0] == ["Date", "Transaction ID", "Description", "Amount", "File Name", "Status", "Notes"]
    assert len(rows) == 1 + 6
    # Every line is there, in statement order (by date), with its status and note.
    assert [row[2] for row in rows[1:]] == [
        "CLICK TO PAY PAYMENT RECEIVED", "MACFOS LIMITED MUMBAI", "MACFOS LIMITED MUMBAI",
        "AMAZON SELLER SERVICES MUMBAI", "FATHOM VIDEO  [36 USD]", "VISTAPRINT INDIA",
    ]
    files = {row[2]: row[4] for row in rows[1:]}
    assert files["AMAZON SELLER SERVICES MUMBAI"] == "200626_Amazon_52390.pdf"  # a Review shows its suggested invoice
    assert files["FATHOM VIDEO  [36 USD]"] == "Fathom_36.pdf" and files["VISTAPRINT INDIA"] in (None, "")
    assert [row[3] for row in rows[1:]][0] == pytest.approx(-60472.0)  # the payment is negative
    assert w.row("L1")["status"] == "Matched"


def test_a_line_with_several_invoices_lists_every_file_sorted_by_name(api):
    w = api.world
    resolve = lambda **body: api.post(f"/api/transactions/{w.line['L6']['id']}/resolve", json=body)  # noqa: E731
    assert resolve(attach_invoice_ids=[w.inv["U1"]["id"], w.inv["V"]["id"]], note="two together").status_code == 200
    _, rows = sheet_rows(download(api), "June 2026")
    assert next(r for r in rows if r[2] == "VISTAPRINT INDIA")[4] == "240626_Amazon_175.pdf, Vistaprint_900.pdf"


def test_the_unused_sheet_matches_the_unused_list_and_leaves_out_what_a_suggestion_holds(api):
    _, rows = sheet_rows(download(api), "Invoices not used")
    assert rows[0] == ["File", "Vendor", "Date", "Amount", "Reason"]
    names = sorted(r[0] for r in rows[1:])
    assert names == ["240626_Amazon_175.pdf", "290626_Amazon_432.pdf", "Vistaprint_900.pdf", "credit_note.pdf", "vendorless.png"]
    assert "200626_Amazon_52390.pdf" not in names  # reserved by L2's Review suggestion (rule 8)
    reasons = {r[0]: r[4] for r in rows[1:]}
    assert reasons["290626_Amazon_432.pdf"] == "Dated after the statement date"


def test_the_summary_sheet_holds_the_statement_figures(api):
    wb, _ = sheet_rows(download(api), "June 2026")
    ws = wb["Summary"]
    printed = [ws.cell(r, 3).value for r in range(14, 22) if ws.cell(r, 1).value in ("Debits (Purchases/Charges)", "Credits (Payments/Credits)")]
    assert [float(p) for p in printed] == [DEBITS / 100, CREDITS / 100]


# ---- exporting, editing, exporting again -------------------------------------------------------------------------------------------
def test_the_first_export_sets_exported_and_records_it(api):
    w = api.world
    assert cycle(api)["status"] == "matched" and cycle(api)["edited_since_export"] is False
    download(api)
    c = cycle(api)
    assert c["status"] == "exported" and c["edited_since_export"] is False and c["last_exported_at"] is not None
    events = [e for e in w.repo.events if e["action"] == "exported"]
    assert len(events) == 1 and events[0]["payload"]["file"] == "CC Expense June 2026.xlsx" and events[0]["payload"]["lines"] == 6


def test_exporting_again_is_always_allowed_and_changes_nothing_else(api):
    download(api)
    download(api)
    assert cycle(api)["status"] == "exported"
    assert [e["action"] for e in api.world.repo.events].count("exported") == 2
    assert "edited_after_export" not in [e["action"] for e in api.world.repo.events]


def test_the_first_edit_after_an_export_moves_the_cycle_back_with_an_edited_since_export_note(api):
    w = api.world
    download(api)
    r = api.post(f"/api/transactions/{w.line['L3']['id']}/resolve", json={"status": "No invoice", "note": "asked vendor"})
    assert r.status_code == 200
    c = cycle(api)
    assert c["status"] == "matched" and c["edited_since_export"] is True and c["last_exported_at"] is not None
    s = summary(api)
    assert s["edited_since_export"] is True and s["last_exported_at"] is not None
    assert [e["action"] for e in w.repo.events].count("edited_after_export") == 1

    # A second edit does not repeat the move.
    api.post(f"/api/transactions/{w.line['L6']['id']}/resolve", json={"note": "chasing"})
    assert [e["action"] for e in w.repo.events].count("edited_after_export") == 1


def test_exporting_again_after_an_edit_sets_exported_and_the_workbook_shows_the_edit(api):
    w = api.world
    download(api)
    api.post(f"/api/transactions/{w.line['L3']['id']}/resolve", json={"status": "No invoice", "note": "asked vendor"})
    r = download(api)
    c = cycle(api)
    assert c["status"] == "exported" and c["edited_since_export"] is False
    _, rows = sheet_rows(r, "June 2026")
    edited = [row for row in rows if row[5] == "No invoice"]
    assert len(edited) == 1 and edited[0][6] == "asked vendor"
    wb = load_workbook(BytesIO(r.content))
    assert wb["Summary"]["A6"].value == "Review"  # the counts are formulas, so they follow the sheet


def test_a_preview_or_a_refused_change_is_not_an_edit(api):
    w = api.world
    download(api)
    api.post(f"/api/transactions/{w.line['L3']['id']}/resolve", json={"status": "No invoice", "preview": True})
    api.post(f"/api/transactions/{w.line['L3']['id']}/resolve", json={"status": "Payment"})  # refused (422)
    assert cycle(api)["status"] == "exported" and cycle(api)["edited_since_export"] is False


def test_a_cycle_that_was_never_exported_is_never_edited_since_export(api):
    w = api.world
    api.post(f"/api/transactions/{w.line['L3']['id']}/resolve", json={"status": "No invoice", "note": "x"})
    c = cycle(api)
    assert c["status"] == "matched" and c["edited_since_export"] is False and c["last_exported_at"] is None
    assert "edited_after_export" not in [e["action"] for e in w.repo.events]
