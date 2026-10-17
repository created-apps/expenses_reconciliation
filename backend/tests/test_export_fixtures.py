"""The export on the real June cycle (Docker only): 127 rows that agree with the expected CSV, the totals check
of test 11.4, the eight Amazon files of test 11.3, and the unused-invoices sheet. Nothing is uploaded anywhere."""

import csv
import shutil
import sys
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from june_cycle import JUNE_ZIP, STATEMENT, build_june
from app.services.export import build_export
from app.services.run_matching import run_matching, start_matching
from test_excel_engine import evaluate

DOC_CANDIDATES = [Path(__file__).resolve().parents[2] / "doc" / "reference", Path("/doc/reference")]
EXPECTED_CSV = next((p / "june_2026_expected_results.csv" for p in DOC_CANDIDATES if (p / "june_2026_expected_results.csv").exists()), None)
SAME_FILE = {"140626_Amazon_525.pdf": "140626_Amazon_525(1).pdf", "MH2727WI016880_GC001_10-Jun-2026.pdf": "DTDC_488.pdf"}

pytestmark = [
    pytest.mark.skipif(not (STATEMENT and JUNE_ZIP and EXPECTED_CSV), reason="June fixtures or expected CSV missing"),
    pytest.mark.skipif(shutil.which("pdftotext") is None, reason="pdftotext (poppler) not installed; run in Docker"),
]


def canon(name: str) -> str:
    return SAME_FILE.get(name, name)


@pytest.fixture(scope="module")
def june():
    j = build_june()
    cycle = j.repo.get_cycle(j.cid)
    start_matching(j.repo, cycle)
    assert run_matching(j.repo, cycle, "t") is not None
    content, file_name = build_export(j.repo, j.repo.get_cycle(j.cid), "t")
    j.content, j.file_name = content, file_name
    j.wb = load_workbook(BytesIO(content))
    return j


@pytest.fixture(scope="module")
def expected():
    with open(EXPECTED_CSV, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def day(text: str) -> date:
    d, m, y = text.split("/")
    return date(int(y), int(m), int(d))


def test_the_workbook_has_three_sheets_and_is_named_for_the_month(june):
    assert june.wb.sheetnames == ["Summary", "June 2026", "Invoices not used"]
    assert june.file_name == "CC Expense June 2026.xlsx"
    assert june.repo.get_cycle(june.cid)["status"] == "exported"


def test_the_main_sheet_has_127_rows_that_agree_with_the_expected_results(june, expected):
    ws = june.wb["June 2026"]
    assert [c.value for c in ws[1]] == ["Date", "Transaction ID", "Description", "Amount", "File Name", "Status", "Notes"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == len(expected) == 127
    by_serial = {str(r[1]): r for r in rows}
    problems = []
    for e in expected:
        r = by_serial[e["ser_no"]]
        want_files = sorted(canon(n.strip()) for n in e["files"].split(",") if n.strip())
        got_files = sorted(canon(n.strip()) for n in (r[4] or "").split(",") if n.strip())
        same = (
            r[0].date() == day(e["date"]) and r[2] == e["description"]
            and Decimal(str(r[3])) == Decimal(e["amount_inr"]) and r[5] == e["status"] and got_files == want_files
        )
        if not same:
            problems.append((e["date"], e["ser_no"], e["description"], e["status"], want_files, r))
    assert problems == []


def test_every_transaction_id_is_a_number_with_all_its_digits(june):
    ws = june.wb["June 2026"]
    for row in ws.iter_rows(min_row=2):
        assert isinstance(row[1].value, int) and row[1].number_format == "0" and row[1].value > 10**9


def test_11_3_the_amazon_charge_lists_its_eight_files_separated_by_comma_and_space(june):
    ws = june.wb["June 2026"]
    row = next(r for r in ws.iter_rows(min_row=2) if r[1].value == 13598184078)
    names = row[4].value.split(", ")
    assert len(names) == 8 and names == sorted(names)
    assert "130626_Amazon_187.16(2).pdf" in names and "130626_amazon_189(1).pdf" in names


def test_11_4_the_summary_check_block_reconciles_with_the_statement(june):
    ws = june.wb["Summary"]
    rows = {ws.cell(r, 1).value: r for r in range(1, 30) if ws.cell(r, 1).value}
    d, c = rows["Debits (Purchases/Charges)"], rows["Credits (Payments/Credits)"]
    assert evaluate(june.wb, ws, f"B{d}") == Decimal("1316107.04") and Decimal(str(ws.cell(d, 3).value)) == Decimal("1316107.04")
    assert evaluate(june.wb, ws, f"B{c}") == Decimal("591066.42") and Decimal(str(ws.cell(c, 3).value)) == Decimal("591066.42")
    assert evaluate(june.wb, ws, f"D{d}") == 0 and evaluate(june.wb, ws, f"D{c}") == 0


def test_the_summary_counts_are_the_june_counts(june):
    ws = june.wb["Summary"]
    counts = {ws.cell(r, 1).value: evaluate(june.wb, ws, f"B{r}") for r in range(5, 13)}
    assert counts == {"Matched": 65, "Review": 5, "Missing invoice": 20, "No invoice": 6, "Reversed": 8,
                      "Credit – manual": 4, "Payment": 1, "Bank charge": 18}
    amounts = {ws.cell(r, 1).value: evaluate(june.wb, ws, f"C{r}") for r in range(5, 13)}
    assert amounts["Matched"] == Decimal("541726.64") and amounts["Payment"] == Decimal("-60472.00")
    assert amounts["Credit – manual"] == Decimal("-55985.60") and amounts["Reversed"] == 0


def test_11_7_the_unused_sheet_lists_the_9_16_files_that_no_suggestion_holds(june):
    ws = june.wb["Invoices not used"]
    assert [c.value for c in ws[1]] == ["File", "Vendor", "Date", "Amount", "Reason"]
    names = [r[0].value for r in ws.iter_rows(min_row=2)]  # one row per invoice: a PDF with two unused invoices appears twice
    for wanted in ("240626_Amazon_175.pdf", "040626_Sembly_USD29.pdf", "created invc.pdf", "CREATED-25947.pdf",
                   "Fathom Video$20.pdf", "Anthropic.pdf", "300626_Macfos_801.pdf"):
        assert wanted in names, wanted
    assert all(any(n.startswith(p) for n in names) for p in ("280626_Amazon_", "290626_", "300626_"))
    reasons = {r[0].value: r[4].value for r in ws.iter_rows(min_row=2)}
    assert reasons["300626_Macfos_801.pdf"] == "Dated after the statement date"
    assert all(reasons.values())  # every row says why


def test_a_review_suggestion_keeps_its_invoices_out_of_the_unused_sheet(june):
    ws = june.wb["Invoices not used"]
    names = {r[0].value for r in ws.iter_rows(min_row=2)}
    main = june.wb["June 2026"]
    suggested = set()
    for r in main.iter_rows(min_row=2):
        if r[5].value == "Review":
            suggested |= {canon(n.strip()) for n in (r[4].value or "").split(",") if n.strip()}
    assert suggested and not {canon(n) for n in names} & suggested


def test_the_new_unused_reasons_on_june(june):
    ws = june.wb["Invoices not used"]
    rows = [(r[0].value, r[4].value) for r in ws.iter_rows(min_row=2)]
    reasons = dict(rows)
    after_last = "Dated after the last Amazon charge on this statement (27/06)"  # the last Amazon debit is 27/06
    assert sum(1 for _, why in rows if why == after_last) == 11  # the Amazon invoices dated 28/06
    assert reasons["280626_Amazon_57.pdf"] == after_last and reasons["280226_Amazon_489.pdf"] == after_last
    assert reasons["240626_Amazon_175.pdf"] == "No charge of this amount found"  # dated before the last charge
    assert reasons["280626_Amazon_123(1).pdf"] == "Dated after the statement date"  # its invoice date is 30/06
    assert reasons["290626_Amazon_1599.pdf"] == "Dated after the statement date"
    assert reasons["CREATED-25947.pdf"] == "No charge of this amount found"
    assert reasons["040626_Sembly_USD29.pdf"] == "No statement line for this vendor"
    # Fathom: a $20.00 charge exists (28/05), but 29 days from the invoice, so it was not matched.
    assert reasons["Fathom Video$20.pdf"] == (
        "A charge of $20.00 exists on 28/05/2026, 29 days from the invoice date (invoices are matched within 7 days)"
    )
    assert len(rows) == 22


def test_june_has_no_amazon_report_so_there_is_no_fetch_list_and_no_amazon_invoice_needs_an_order_id(june):
    from app.services.amazon_fetch import amazon_to_fetch

    cycle = june.repo.get_cycle(june.cid)
    view = amazon_to_fetch(june.repo, cycle)
    assert view["has_report"] is False and view["items"] == [] and view["coverage"] is None
    amazon = [i for i in june.repo.list_invoices(june.cid) if i["vendor_key"] == "amazon"]
    assert len(amazon) == 64 and all(i["extraction_status"] == "ok" and i["order_id"] for i in amazon)
    bad = [i for i in june.repo.list_invoices(june.cid) if i["extraction_status"] == "needs_check"]
    assert len(bad) == 1  # only the Mailchimp receipt (no total line), exactly as before the Order ID rule

# ---- the read-only comparison script (scripts/compare_unused_sheet.py) ------------------------------------------------
def load_script():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import compare_unused_sheet

    return compare_unused_sheet


def reference_workbook(files: list[str]) -> bytes:
    wb = Workbook()
    wb.active.title = "CC Expense"
    ws = wb.create_sheet("Invoices not used")
    ws.append(["Reason", "File"])  # the file column need not be the first
    for f in files:
        ws.append(["x", f])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_the_comparison_script_reads_the_file_column_and_reports_both_directions():
    script = load_script()
    names = script.reference_names(reference_workbook(["a.pdf", "b.pdf", "MH2727WI016880_GC001_10-Jun-2026.pdf"]))
    assert names == ["a.pdf", "b.pdf", "MH2727WI016880_GC001_10-Jun-2026.pdf"]
    only_app, only_ref = script.differences({"a.pdf": "r", "DTDC_488.pdf": "r", "c.pdf": "r"}, names)
    assert only_app == ["c.pdf"] and only_ref == ["b.pdf"]  # byte-identical pairs count as the same file


def test_the_comparison_script_on_the_june_export_itself_finds_no_difference(june):
    script = load_script()
    ours = {r[0].value: r[4].value for r in june.wb["Invoices not used"].iter_rows(min_row=2)}
    assert script.differences(ours, script.reference_names(june.content)) == ([], [])
