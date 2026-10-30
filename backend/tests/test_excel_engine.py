"""The workbook layout of project_context.md section 7 (engine/excel.py), read back with openpyxl."""

import re
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.engine.excel import (
    HEADERS,
    INR_FORMAT,
    STATUSES,
    SheetLine,
    Statement,
    UnusedRow,
    build_workbook,
    check_totals,
    export_file_name,
    line_description,
    sheet_name_for,
)


def line(ser="13598184078", day=date(2026, 6, 13), desc="ASSPL Bangalore IN", paise=255744, credit=False,
         files=(), status="Matched", note="", foreign=None):
    return SheetLine(day, ser, desc, paise, credit, foreign[0] if foreign else None, foreign[1] if foreign else None,
                     tuple(files), status, note)


EIGHT = tuple(f"130626_Amazon_{n}.pdf" for n in ("187.16(2)", "198", "220", "238", "499", "928.27", "98.01")) + ("130626_amazon_189(1).pdf",)

LINES = [
    line(files=EIGHT, note="8 invoices sum exactly to the charge"),
    line("13576983915", date(2026, 6, 11), "Late Payment Fee", 130000, status="Bank charge", note="Card fee"),
    line("13558105225", date(2026, 6, 6), "SCRIPTIT.FR/CANCEL PARIS FR*", 89669, files=(), status="Review", foreign=(899, "USD")),
    line("13568000001", date(2026, 6, 10), "CLICK TO PAY PAYMENT RECEIVED", 6047200, credit=True, status="Payment"),
]
STATEMENT = Statement(debits_paise=255744 + 130000 + 89669, credits_paise=6047200, period_start=date(2026, 5, 21), period_end=date(2026, 6, 28))
UNUSED = [
    UnusedRow("240626_Amazon_175.pdf", "Amazon", date(2026, 6, 24), "INR", 17500, "No charge of this amount found"),
    UnusedRow("040626_Sembly_USD29.pdf", "Sembly", date(2026, 6, 4), "USD", 2900, "No statement line for this vendor"),
    UnusedRow("vendorless.png", "", None, "INR", None, "No vendor set. Open the invoice and choose one"),
]


def workbook(lines=LINES, unused=UNUSED, statement=STATEMENT, sheet="June 2026"):
    return load_workbook(BytesIO(build_workbook(sheet, lines, unused, statement)))


# ---- sheets and the main sheet ------------------------------------------------------------------------------------------------
def test_three_sheets_in_the_order_of_the_spec():
    assert workbook().sheetnames == ["Summary", "June 2026", "Invoices not used"]


def test_the_main_sheet_headers_are_exactly_the_clients():
    ws = workbook()["June 2026"]
    assert [c.value for c in ws[1]] == ["Date", "Transaction ID", "Description", "Amount", "File Name", "Status", "Notes"]
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.max_row == 1 + len(LINES)


def test_dates_are_real_dates_in_d_m_yyyy():
    ws = workbook()["June 2026"]
    assert ws["A2"].value.date() == date(2026, 6, 13) and ws["A2"].number_format == "d/m/yyyy"


def test_the_transaction_id_is_a_number_written_with_every_digit():
    ws = workbook()["June 2026"]
    assert ws["B2"].value == 13598184078 and isinstance(ws["B2"].value, int)
    assert ws["B2"].number_format == "0"  # shows 13598184078, never 1.36E+10


def test_foreign_lines_append_the_foreign_amount_to_the_description():
    ws = workbook()["June 2026"]
    assert ws["C4"].value == "SCRIPTIT.FR/CANCEL PARIS FR*  [8.99 USD]"
    assert ws["C2"].value == "ASSPL Bangalore IN"
    assert line_description("X", 8612, "USD") == "X  [86.12 USD]" and line_description("X", 5, "USD") == "X  [0.05 USD]"
    # Round amounts are written short, as in the client's earlier workbooks.
    assert line_description("X", 3600, "USD") == "X  [36 USD]" and line_description("X", 13350, "USD") == "X  [133.5 USD]"
    assert line_description("X", None, None) == "X"


def test_amounts_are_numbers_with_credits_negative_and_the_rupee_format():
    ws = workbook()["June 2026"]
    assert ws["D2"].value == Decimal("2557.44") or float(ws["D2"].value) == 2557.44
    assert float(ws["D5"].value) == -60472.0  # the payment is a credit
    for r in range(2, 6):
        assert ws.cell(r, 4).data_type == "n" and ws.cell(r, 4).number_format == INR_FORMAT
    assert INR_FORMAT.replace('"', "") == "₹#,##0.00;-₹#,##0.00"  # the client's format (the symbol is only quoted)


def test_file_names_are_separated_by_comma_and_space_and_status_and_notes_follow():
    ws = workbook()["June 2026"]
    assert ws["E2"].value == ", ".join(EIGHT) and ws["E2"].value.count(", ") == 7
    assert ws["E3"].value == "" or ws["E3"].value is None
    assert (ws["F2"].value, ws["G2"].value) == ("Matched", "8 invoices sum exactly to the charge")
    assert ws["F5"].value == "Payment"


def test_everything_is_arial_10():
    wb = workbook()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    assert (cell.font.name, cell.font.sz) == ("Arial", 10), (ws.title, cell.coordinate)


# ---- Invoices not used --------------------------------------------------------------------------------------------------------
def test_the_unused_sheet_lists_file_vendor_date_amount_and_reason():
    ws = workbook()["Invoices not used"]
    assert [c.value for c in ws[1]] == ["File", "Vendor", "Date", "Amount", "Reason"]
    assert [ws.cell(r, 1).value for r in (2, 3, 4)] == ["240626_Amazon_175.pdf", "040626_Sembly_USD29.pdf", "vendorless.png"]
    assert float(ws["D2"].value) == 175.0 and ws["D2"].number_format == INR_FORMAT
    assert float(ws["D3"].value) == 29.0 and "$" in ws["D3"].number_format  # a USD invoice keeps its own currency
    assert ws["D4"].value is None and ws["C4"].value is None  # nothing was read: blank, not zero
    assert ws["E3"].value == "No statement line for this vendor" and ws["B2"].value == "Amazon"


# ---- Summary: formulas and the check block ------------------------------------------------------------------------------------
def evaluate(wb, ws, ref: str):
    """Just enough of Excel for the formulas the Summary uses, so the test proves what they would show."""
    value = ws[ref].value
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    main = wb["June 2026"]

    def rng(text: str):
        col = re.search(r"\$([A-Z])\$2:\$[A-Z]\$(\d+)", text)
        return [main[f"{col.group(1)}{r}"].value for r in range(2, int(col.group(2)) + 1)]

    num = lambda v: Decimal(str(v))  # noqa: E731
    if m := re.fullmatch(r"=COUNTIF\((.+?),(A\d+)\)", value):
        return sum(1 for v in rng(m.group(1)) if v == ws[m.group(2)].value)
    if m := re.fullmatch(r"=SUMIF\((.+?),(A\d+),(.+?)\)", value):
        keys, nums = rng(m.group(1)), rng(m.group(3))
        return sum((num(n) for k, n in zip(keys, nums) if k == ws[m.group(2)].value), Decimal(0))
    if m := re.fullmatch(r'=(-?)SUMIF\((.+?),"([<>])0"\)', value):
        nums = [num(n) for n in rng(m.group(2))]
        total = sum((n for n in nums if (n > 0 if m.group(3) == ">" else n < 0)), Decimal(0))
        return -total if m.group(1) else total
    if m := re.fullmatch(r"=SUM\(([A-Z])(\d+):([A-Z])(\d+)\)", value):
        return sum((num(evaluate(wb, ws, f"{m.group(1)}{r}")) for r in range(int(m.group(2)), int(m.group(4)) + 1)), Decimal(0))
    if m := re.fullmatch(r"=ROUND\(([A-Z]\d+)-([A-Z]\d+),2\)", value):
        return round(num(evaluate(wb, ws, m.group(1))) - num(evaluate(wb, ws, m.group(2))), 2)
    raise AssertionError(f"not a formula this test understands: {value}")


def test_the_summary_counts_and_amounts_per_status_are_formulas_that_add_up_to_the_sheet():
    wb = workbook()
    ws = wb["Summary"]
    assert [ws.cell(r, 1).value for r in range(5, 13)] == list(STATUSES)
    for r in range(5, 13):
        assert str(ws.cell(r, 2).value).startswith("=COUNTIF(") and str(ws.cell(r, 3).value).startswith("=SUMIF(")
    counts = {ws.cell(r, 1).value: evaluate(wb, ws, f"B{r}") for r in range(5, 13)}
    assert counts == {"Matched": 1, "Review": 1, "Missing invoice": 0, "No invoice": 0, "Reversed": 0,
                      "Credit – manual": 0, "Payment": 1, "Bank charge": 1}
    amounts = {ws.cell(r, 1).value: evaluate(wb, ws, f"C{r}") for r in range(5, 13)}
    assert amounts["Matched"] == Decimal("2557.44") and amounts["Payment"] == Decimal("-60472.00")
    assert evaluate(wb, ws, "B13") == 4 and evaluate(wb, ws, "C13") == Decimal("2557.44") + 1300 + Decimal("896.69") - 60472


def test_the_check_block_shows_zero_difference_when_the_sheet_adds_up_to_the_statement():
    wb = workbook()
    ws = wb["Summary"]
    labels = {ws.cell(r, 1).value: r for r in range(14, 22) if ws.cell(r, 1).value}
    d, c = labels["Debits (Purchases/Charges)"], labels["Credits (Payments/Credits)"]
    assert evaluate(wb, ws, f"B{d}") == Decimal("4754.13") and Decimal(str(ws.cell(d, 3).value)) == Decimal("4754.13")
    assert evaluate(wb, ws, f"D{d}") == 0 and evaluate(wb, ws, f"D{c}") == 0
    assert evaluate(wb, ws, f"B{c}") == Decimal("60472.00")
    assert "OK" in str(ws.cell(c + 1, 1).value) and str(ws.cell(c + 1, 1).value).startswith("=IF(AND(")


def test_the_check_block_shows_a_difference_when_it_does_not_add_up():
    off = Statement(debits_paise=255744 + 130000 + 89669 + 100, credits_paise=6047200)
    wb = workbook(statement=off)
    ws = wb["Summary"]
    r = next(r for r in range(14, 22) if ws.cell(r, 1).value == "Debits (Purchases/Charges)")
    assert evaluate(wb, ws, f"D{r}") == Decimal("-1.00")  # the sheet has ₹1.00 less than the statement prints


def test_figures_that_were_not_read_are_marked_not_read():
    wb = workbook(statement=Statement(debits_paise=None, credits_paise=None))
    ws = wb["Summary"]
    r = next(r for r in range(14, 22) if ws.cell(r, 1).value == "Debits (Purchases/Charges)")
    assert (ws.cell(r, 3).value, ws.cell(r, 4).value) == ("not read", "n/a")


def test_an_empty_statement_still_gives_a_valid_workbook():
    wb = workbook(lines=[], unused=[])
    assert wb.sheetnames == ["Summary", "June 2026", "Invoices not used"]
    assert wb["June 2026"].max_row == 1


# ---- the totals check, names ---------------------------------------------------------------------------------------------------------
def test_check_totals_compares_debits_and_credits_with_the_printed_figures():
    t = check_totals(LINES, STATEMENT)
    assert (t.debits_paise, t.credits_paise, t.debits_difference_paise, t.credits_difference_paise, t.ok) == (475413, 6047200, 0, 0, True)
    off = check_totals(LINES, Statement(debits_paise=475413 + 500, credits_paise=6047200))
    assert off.debits_difference_paise == -500 and off.ok is False
    unknown = check_totals(LINES, Statement(debits_paise=None, credits_paise=None))
    assert unknown.ok is False and unknown.debits_difference_paise is None


def test_the_month_sheet_name_comes_from_the_statement_date():
    assert sheet_name_for(date(2026, 6, 28), "whatever") == "June 2026"
    assert sheet_name_for(None, "My label") == "My label"
    assert sheet_name_for(None, "a/b[c]:d?") == "abcd"  # characters Excel forbids in a sheet name
    assert sheet_name_for(None, None) == "Statement" and len(sheet_name_for(None, "x" * 80)) == 31


def test_the_file_name_follows_the_clients():
    assert export_file_name("June 2026") == "CC Expense June 2026.xlsx"
    assert export_file_name("") == "CC Expense.xlsx"
